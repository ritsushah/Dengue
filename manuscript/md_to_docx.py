#!/usr/bin/env python3
"""Convert the combined manuscript Markdown to an editable Word document."""

from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "FIA_vs_ELISA_Dengue_NS1_MetaAnalysis.md"
DOCX_PATH = ROOT / "FIA_vs_ELISA_Dengue_NS1_MetaAnalysis.docx"
DOC_PATH = ROOT / "FIA_vs_ELISA_Dengue_NS1_MetaAnalysis.doc"


def strip_yaml(text: str) -> str:
    if text.startswith("---"):
        match = re.match(r"^---\n.*?\n---\n+", text, flags=re.S)
        if match:
            return text[match.end() :]
    return text


def set_run_font(run, name="Times New Roman", size=11, bold=False, italic=False, code=False):
    if code:
        name = "Courier New"
        size = 9
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if code:
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def add_formatted_runs(paragraph, text: str, base_size=11, base_italic=False):
    pattern = re.compile(
        r"(\*\*[^*]+?\*\*|\*[^*\n]+?\*|`[^`]+?`|\[[^\]]+\]\([^)]+\))"
    )
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=base_size, italic=base_italic)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, italic=base_italic)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, code=True)
        elif token.startswith("[") and "](" in token:
            label, _url = token[1:-1].split("](", 1)
            run = paragraph.add_run(label)
            set_run_font(run, size=base_size, italic=True)
            run.font.color.rgb = RGBColor(0x0B, 0x3D, 0x91)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, italic=base_italic)


def is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped.startswith("|"):
        return False
    inner = stripped.strip("|")
    cells = [c.strip() for c in inner.split("|")]
    return all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells) and len(cells) >= 1


def parse_table_row(line: str) -> list[str]:
    inner = line.strip().strip("|")
    return [c.strip() for c in inner.split("|")]


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            value = row[j] if j < len(row) else ""
            add_formatted_runs(p, value, base_size=8)
            for run in p.runs:
                run.font.size = Pt(8)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for level, size in ((1, 16), (2, 13), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        if level == 3:
            style.font.italic = True
        style.paragraph_format.space_before = Pt(14 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(8)


def convert() -> None:
    text = strip_yaml(MD_PATH.read_text(encoding="utf-8"))
    lines = text.splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    configure_styles(doc)

    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                block = "\n".join(code_lines) if code_lines else ""
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.15)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(10)
                run = p.add_run(block)
                set_run_font(run, code=True, size=8)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        img = re.match(r"^!\[(.*?)\]\((.*?)\)$", line.strip())
        if img:
            caption, rel = img.group(1), img.group(2)
            img_path = (ROOT / rel).resolve()
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run()
                run.add_picture(str(img_path), width=Inches(6.3))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_runs(cap, f"*{caption}*" if not caption.startswith("*") else caption, base_size=9, base_italic=True)
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [parse_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if line.strip() in {"---", "***"}:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("—" * 24)
            set_run_font(run, size=9)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            i += 1
            continue

        bullet = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if bullet:
            body = bullet.group(3)
            numbered = bool(re.match(r"\d+\.", bullet.group(2)))
            style = "List Number" if numbered else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_formatted_runs(p, body)
            i += 1
            continue

        if line.startswith("> "):
            quote = [line[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("> "):
                quote.append(lines[i][2:])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            add_formatted_runs(p, " ".join(quote), base_italic=True)
            continue

        para = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (
                not nxt.strip()
                or nxt.startswith("#")
                or nxt.startswith("```")
                or nxt.startswith("|")
                or nxt.startswith(">")
                or nxt.startswith("![")
                or nxt.strip() in {"---", "***"}
                or re.match(r"^(\s*)([-*]|\d+\.)\s+", nxt)
            ):
                break
            para.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_runs(p, " ".join(x.strip() for x in para))

    doc.save(DOCX_PATH)

    # Native .doc via macOS textutil so Windows Word 97–2003/compatibility mode can open it.
    subprocess.run(
        ["textutil", "-convert", "doc", "-output", str(DOC_PATH), str(DOCX_PATH)],
        check=True,
    )


if __name__ == "__main__":
    convert()
    print(f"Wrote {DOCX_PATH.name} ({DOCX_PATH.stat().st_size / 1024:.0f} KB)")
    print(f"Wrote {DOC_PATH.name} ({DOC_PATH.stat().st_size / 1024:.0f} KB)")
